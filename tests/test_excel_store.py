from datetime import date, datetime

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import excel_store
from conftest import make_day, make_days
from tnb_data import COLUMNS


def legacy_workbook(path, rows, duplicate_last=False):
    """Old 3-column file, as written by the previous app."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Usage"
    ws.append(["Timestamp", "kWh", "RatePlan"])
    for r in rows:
        ws.append([r.timestamp, r.kwh, r.rateplan])
    if duplicate_last:
        for r in rows[-48:]:
            ws.append([r.timestamp, r.kwh, r.rateplan])
    wb.save(path)


def test_account_file_name(tmp_path):
    assert excel_store.account_file(tmp_path, "220000000001") == tmp_path / "tnb_usage_220000000001.xlsx"


def test_new_file_has_four_columns_and_rows(tmp_path):
    path = tmp_path / "a.xlsx"
    assert excel_store.append_rows(path, make_day(date(2026, 10, 1))) == 48
    ws = load_workbook(path)["Usage"]
    assert [c.value for c in ws[1]] == list(COLUMNS)
    assert ws.cell(row=2, column=1).value == datetime(2026, 10, 1, 0, 0)
    assert ws.cell(row=2, column=4).value == "off_peak_energy_charge"


def test_append_skips_duplicates(tmp_path):
    path = tmp_path / "a.xlsx"
    excel_store.append_rows(path, make_days(date(2026, 10, 1), date(2026, 10, 3)))
    assert excel_store.append_rows(path, make_days(date(2026, 10, 2), date(2026, 10, 4))) == 48
    assert excel_store.append_rows(path, make_days(date(2026, 10, 1), date(2026, 10, 4))) == 0
    df = excel_store.read_frame(path)
    assert len(df) == 3 * 48
    assert df["Timestamp"].is_unique


def test_append_empty_is_noop(tmp_path):
    path = tmp_path / "a.xlsx"
    assert excel_store.append_rows(path, []) == 0
    assert not path.exists()


def test_watermark(tmp_path):
    path = tmp_path / "a.xlsx"
    assert excel_store.watermark(path) is None
    excel_store.append_rows(path, make_day(date(2026, 10, 1))[:10])
    assert excel_store.watermark(path) == datetime(2026, 10, 1, 4, 30)


def test_legacy_file_gets_category_header_on_append(tmp_path):
    path = tmp_path / "old.xlsx"
    legacy_workbook(path, make_day(date(2026, 8, 12), category=""))
    excel_store.append_rows(path, make_day(date(2026, 8, 13)))
    ws = load_workbook(path)["Usage"]
    assert [c.value for c in ws[1]] == list(COLUMNS)
    df = excel_store.read_frame(path)
    assert len(df) == 96
    assert df.iloc[0]["Category"] == "" and df.iloc[-1]["Category"] == "off_peak_energy_charge"


def test_backfill_keeps_sheet_sorted(tmp_path):
    path = tmp_path / "a.xlsx"
    excel_store.append_rows(path, make_day(date(2026, 1, 1)) + make_day(date(2026, 1, 3)))
    assert excel_store.find_gaps(excel_store.read_frame(path)) == [date(2026, 1, 2)]
    assert excel_store.append_rows(path, make_day(date(2026, 1, 2), kwh=0.9)) == 48
    ws = load_workbook(path)["Usage"]
    stamps = [r[0] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True)]
    assert stamps == sorted(stamps) and len(stamps) == 144
    assert excel_store.find_gaps(excel_store.read_frame(path)) == []


def test_find_gaps_ignores_partial_last_day():
    rows = make_days(date(2026, 1, 1), date(2026, 1, 3)) + make_day(date(2026, 1, 3), points=5)
    rows = [r for r in rows if r.timestamp != datetime(2026, 1, 2, 12, 0)]
    df = pd.DataFrame({"Timestamp": [r.timestamp for r in rows], "kWh": [r.kwh for r in rows]})
    assert excel_store.find_gaps(df) == [date(2026, 1, 2)]
    assert excel_store.find_gaps(excel_store._empty_frame()) == []


def test_locked_file_raises_friendly_error(tmp_path, monkeypatch):
    path = tmp_path / "a.xlsx"
    excel_store.append_rows(path, make_day(date(2026, 1, 1)))

    def locked(*_a, **_k):
        raise PermissionError("in use")

    monkeypatch.setattr(excel_store.os, "replace", locked)
    with pytest.raises(excel_store.ExcelLockedError, match="close it in Excel"):
        excel_store.append_rows(path, make_day(date(2026, 1, 2)))
    assert len(excel_store.read_frame(path)) == 48
    assert not list(tmp_path.glob("*.saving.xlsx"))


def test_migrate_legacy_backs_up_dedupes_and_moves(tmp_path):
    legacy = tmp_path / "tnb_usage.xlsx"
    rows = make_days(date(2026, 8, 10), date(2026, 8, 13), category="")
    legacy_workbook(legacy, rows, duplicate_last=True)
    original_bytes = legacy.read_bytes()

    target = excel_store.migrate_legacy(legacy, tmp_path, "220000000001", today=date(2026, 10, 28))

    assert target == tmp_path / "tnb_usage_220000000001.xlsx"
    backup = tmp_path / "tnb_usage_backup_20261028.xlsx"
    assert backup.read_bytes() == original_bytes
    assert not legacy.exists()
    df = excel_store.read_frame(target)
    assert len(df) == 3 * 48 and df["Timestamp"].is_unique
    assert excel_store.watermark(target) == datetime(2026, 8, 12, 23, 30)
    assert [c.value for c in load_workbook(target)["Usage"][1]] == list(COLUMNS)


def test_migrate_legacy_does_nothing_when_target_exists(tmp_path):
    legacy = tmp_path / "tnb_usage.xlsx"
    legacy_workbook(legacy, make_day(date(2026, 8, 12)))
    excel_store.append_rows(excel_store.account_file(tmp_path, "1"), make_day(date(2026, 8, 12)))
    assert excel_store.migrate_legacy(legacy, tmp_path, "1") is None
    assert legacy.exists()
    assert excel_store.migrate_legacy(tmp_path / "missing.xlsx", tmp_path, "2") is None
