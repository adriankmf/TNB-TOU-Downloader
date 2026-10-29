"""
excel_store.py
--------------
One Excel workbook per myTNB account: <folder>/tnb_usage_<caNo>.xlsx, sheet "Usage",
columns Timestamp | kWh | RatePlan | Category.

  * watermark()      -> last Timestamp saved (where the next sync starts)
  * append_rows()    -> add only new timestamps (never duplicates); a back-filled gap is
                        merged in date order
  * find_gaps()      -> days inside the saved history that have fewer than 48 points
  * migrate_legacy() -> turn the old single tnb_usage.xlsx into one account's file (backup first)
"""

from __future__ import annotations

import logging
import os
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook, load_workbook

from tnb_data import (COL_CATEGORY, COL_KWH, COL_RATEPLAN, COL_TIMESTAMP, COLUMNS,
                      POINTS_PER_DAY, UsageRow, days_between)

logger = logging.getLogger(__name__)

SHEET_NAME = "Usage"
TIMESTAMP_FORMAT = "yyyy-mm-dd hh:mm:ss"


class ExcelLockedError(RuntimeError):
    """The workbook could not be written -- usually because it is open in Excel."""


def account_file(folder: str | Path, ca_no: str) -> Path:
    return Path(folder) / f"tnb_usage_{ca_no}.xlsx"


# -- Reading ------------------------------------------------------------------------------------

def read_frame(path: Path) -> pd.DataFrame:
    """All rows as a DataFrame with the four standard columns (sorted, de-duplicated)."""
    if not path.exists():
        return _empty_frame()
    df = pd.read_excel(path, sheet_name=SHEET_NAME)
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df = df[list(COLUMNS)]
    df[COL_TIMESTAMP] = pd.to_datetime(df[COL_TIMESTAMP], errors="coerce")
    df[COL_KWH] = pd.to_numeric(df[COL_KWH], errors="coerce")
    df[COL_RATEPLAN] = df[COL_RATEPLAN].fillna("").astype(str)
    df[COL_CATEGORY] = df[COL_CATEGORY].fillna("").astype(str)
    df = df.dropna(subset=[COL_TIMESTAMP])
    return df.drop_duplicates(COL_TIMESTAMP).sort_values(COL_TIMESTAMP).reset_index(drop=True)


def watermark(path: Path) -> datetime | None:
    """Latest Timestamp in the file, or None when the file is missing/empty."""
    if not path.exists():
        return None
    wb = load_workbook(path, read_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            return None
        stamps = [v for (v,) in wb[SHEET_NAME].iter_rows(min_row=2, max_col=1, values_only=True)
                  if isinstance(v, datetime)]
    finally:
        wb.close()
    return max(stamps) if stamps else None


def find_gaps(df: pd.DataFrame) -> list[date]:
    """Days between the first and last saved day that have fewer than 48 points."""
    if df.empty:
        return []
    counts = df.groupby(df[COL_TIMESTAMP].dt.date).size()
    first, last = counts.index.min(), counts.index.max()
    # The last day may still be filling in -- it is not a gap.
    return [d for d in days_between(first, last) if counts.get(d, 0) < POINTS_PER_DAY]


# -- Writing ------------------------------------------------------------------------------------

def append_rows(path: Path, rows: Iterable[UsageRow]) -> int:
    """Write rows whose Timestamp is not in the file yet. Returns how many were written."""
    incoming = {r.timestamp: r for r in rows}
    if not incoming:
        return 0

    if not path.exists():
        _write_new(path, sorted(incoming.values(), key=lambda r: r.timestamp))
        return len(incoming)

    wb = load_workbook(path)
    ws = _usage_sheet(wb)
    existing = [row for row in ws.iter_rows(min_row=2, max_col=4, values_only=True)
                if isinstance(row[0], datetime)]
    known = {row[0] for row in existing}
    new_rows = sorted((r for ts, r in incoming.items() if ts not in known),
                      key=lambda r: r.timestamp)
    if not new_rows:
        wb.close()
        return 0

    latest = max(known) if known else None
    if latest is not None and new_rows[0].timestamp < latest:
        # Back-filling a gap: rewrite the data rows in date order so the sheet stays sorted.
        merged = [tuple(_pad(row)) for row in existing] + [_values(r) for r in new_rows]
        merged.sort(key=lambda v: v[0])
        ws.delete_rows(2, ws.max_row)
        for values in merged:
            _append(ws, values)
    else:
        for r in new_rows:
            _append(ws, _values(r))

    _save(wb, path)
    return len(new_rows)


def _usage_sheet(wb):
    if SHEET_NAME not in wb.sheetnames:
        ws = wb.create_sheet(SHEET_NAME)
        ws.append(list(COLUMNS))
        return ws
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    for idx, name in enumerate(COLUMNS, start=1):
        if len(header) < idx or header[idx - 1] != name:
            ws.cell(row=1, column=idx, value=name)
    return ws


def _values(row: UsageRow) -> tuple:
    return (row.timestamp, row.kwh, row.rateplan, row.category)


def _pad(row: tuple) -> list:
    values = list(row) + [None] * (4 - len(row))
    values[2] = values[2] or ""
    values[3] = values[3] or ""
    return values


def _append(ws, values: tuple) -> None:
    ws.append(list(values))
    ws.cell(row=ws.max_row, column=1).number_format = TIMESTAMP_FORMAT


def _write_new(path: Path, rows: list[UsageRow]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(list(COLUMNS))
    for r in rows:
        _append(ws, _values(r))
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 26
    _save(wb, path)


def _save(wb, path: Path) -> None:
    """Save via a temp file so a crash or a locked file never leaves a half-written workbook."""
    tmp = path.with_name(path.stem + ".saving.xlsx")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as exc:
        tmp.unlink(missing_ok=True)
        raise ExcelLockedError(f"Cannot write {path.name} -- close it in Excel and sync again.") from exc
    finally:
        wb.close()


def _empty_frame() -> pd.DataFrame:
    return pd.DataFrame({
        COL_TIMESTAMP: pd.Series(dtype="datetime64[ns]"),
        COL_KWH: pd.Series(dtype=float),
        COL_RATEPLAN: pd.Series(dtype=str),
        COL_CATEGORY: pd.Series(dtype=str),
    })


# -- One-off migration --------------------------------------------------------------------------

def migrate_legacy(legacy: Path, folder: Path, ca_no: str, today: date | None = None) -> Path | None:
    """
    Move the old single-account tnb_usage.xlsx to <folder>/tnb_usage_<caNo>.xlsx.

    A byte-for-byte backup is copied first. The new file is de-duplicated, sorted and gets the
    Category column. The original is removed only after the new file is verified.
    Returns the new path, or None when there is nothing to migrate.
    """
    target = account_file(folder, ca_no)
    if not legacy.exists() or target.exists() or legacy.resolve() == target.resolve():
        return None

    stamp = (today or date.today()).strftime("%Y%m%d")
    backup = legacy.with_name(f"{legacy.stem}_backup_{stamp}{legacy.suffix}")
    shutil.copy2(legacy, backup)
    if backup.stat().st_size != legacy.stat().st_size:
        raise OSError(f"Backup copy of {legacy.name} is incomplete -- migration stopped.")

    df = read_frame(legacy)
    rows = [UsageRow(ts.to_pydatetime(), float(kwh), plan, cat)
            for ts, kwh, plan, cat in df.itertuples(index=False, name=None)]
    _write_new(target, rows)

    if len(read_frame(target)) != len(df):
        raise OSError(f"{target.name} does not match {legacy.name} -- original kept.")
    legacy.unlink()
    logger.info("Migrated %s -> %s (%d rows, backup %s).", legacy.name, target.name, len(df), backup.name)
    return target
